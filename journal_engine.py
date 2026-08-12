from __future__ import annotations

import json
import math
import os
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from dataclasses import dataclass, asdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Flowable, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

PAGE_SIZE = landscape(A4)
PAGE_W, PAGE_H = PAGE_SIZE
LEFT = 10 * mm
RIGHT = 10 * mm
TOP = 10 * mm
BOTTOM = 10 * mm
HEADER_H = 15 * mm


def decode_google_unicode_name(text: str) -> str:
    """Decode Google Drive ZIP names like #U0424#U0456... without touching file contents."""
    return re.sub(r"#U([0-9A-Fa-f]{4})", lambda m: chr(int(m.group(1), 16)), text)


def safe_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def visible_cell_value(formula_cell: Cell, value_cell: Cell) -> str:
    """Use Excel's cached/displayed result for formulas; otherwise preserve the stored value."""
    if formula_cell.data_type == "f":
        cached = value_cell.value
        if cached is None:
            return safe_str(formula_cell.value)
        return safe_str(cached)
    return safe_str(formula_cell.value)


def is_excel_error(value: Any) -> bool:
    return isinstance(value, str) and value.upper() in {
        "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA"
    }


@dataclass
class WarningItem:
    id: str
    kind: str
    message: str
    subject: str = ""
    sheet: str = ""
    cell: str = ""
    original: str = ""
    formula: str = ""
    correctable: bool = False


@dataclass
class JournalSource:
    file_path: str
    display_file: str
    sheet: str
    subject: str
    teacher: str
    group: str
    grade_header_row: Optional[int]
    grade_date_col: Optional[int]
    student_index_col: Optional[int]
    student_name_col: Optional[int]
    student_rows: List[int]
    grade_cols: List[int]
    note_rows: List[int]
    theme_header_row: Optional[int]
    theme_date_col: Optional[int]
    theme_hours_col: Optional[int]
    theme_text_col: Optional[int]
    theme_prefix_cols: List[int]
    theme_rows: List[int]


class PlainCellText(Flowable):
    """Plain-text Flowable that wraps without interpreting markup and without collapsing newlines."""
    def __init__(self, text: str, font_name: str, font_size: float, align: str = "left", leading: Optional[float] = None):
        super().__init__()
        self.text = "" if text is None else str(text)
        self.font_name = font_name
        self.font_size = font_size
        self.leading = leading or font_size * 1.14
        self.align = align
        self.lines: List[str] = []

    def _split_long_token(self, token: str, width: float) -> List[str]:
        if pdfmetrics.stringWidth(token, self.font_name, self.font_size) <= width:
            return [token]
        out, current = [], ""
        for ch in token:
            probe = current + ch
            if current and pdfmetrics.stringWidth(probe, self.font_name, self.font_size) > width:
                out.append(current)
                current = ch
            else:
                current = probe
        if current or not out:
            out.append(current)
        return out

    def _wrap_line(self, raw: str, width: float) -> List[str]:
        if raw == "":
            return [""]
        # Retain spaces as tokens; wrapping may move a space to the next line but does not alter cell data.
        tokens = re.findall(r"\S+|\s+", raw)
        lines: List[str] = []
        current = ""
        for token in tokens:
            pieces = self._split_long_token(token, width)
            for piece in pieces:
                probe = current + piece
                if current and pdfmetrics.stringWidth(probe, self.font_name, self.font_size) > width:
                    lines.append(current.rstrip("\r\n"))
                    current = piece.lstrip("\r\n")
                else:
                    current = probe
        lines.append(current.rstrip("\r\n"))
        return lines

    def wrap(self, availWidth, availHeight):
        width = max(1.0, availWidth)
        self.lines = []
        for raw in self.text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
            self.lines.extend(self._wrap_line(raw, width))
        if not self.lines:
            self.lines = [""]
        self.width = width
        self.height = max(self.leading, len(self.lines) * self.leading)
        return self.width, self.height

    def draw(self):
        self.canv.setFont(self.font_name, self.font_size)
        y = self.height - self.font_size
        for line in self.lines:
            w = pdfmetrics.stringWidth(line, self.font_name, self.font_size)
            if self.align == "center":
                x = max(0, (self.width - w) / 2)
            elif self.align == "right":
                x = max(0, self.width - w)
            else:
                x = 0
            self.canv.drawString(x, y, line)
            y -= self.leading


class VerticalText(Flowable):
    def __init__(self, text: str, font_name: str, font_size: float):
        super().__init__()
        self.text = "" if text is None else str(text)
        self.font_name = font_name
        self.font_size = font_size
        self.text_width = pdfmetrics.stringWidth(self.text, font_name, font_size)

    def wrap(self, availWidth, availHeight):
        self.width = availWidth
        self.height = availHeight
        return availWidth, availHeight

    def draw(self):
        self.canv.saveState()
        self.canv.setFont(self.font_name, self.font_size)
        self.canv.translate(self.width / 2, 2)
        self.canv.rotate(90)
        self.canv.drawString(0, -self.font_size * 0.33, self.text)
        self.canv.restoreState()


class FontBundle:
    def __init__(self, regular: str, bold: str, italic: str, bolditalic: str, label: str):
        self.regular = regular
        self.bold = bold
        self.italic = italic
        self.bolditalic = bolditalic
        self.label = label


def _candidate_font_files() -> List[Tuple[str, List[str]]]:
    # Do not redistribute fonts. The program only uses fonts already installed on the machine.
    return [
        ("Times New Roman", [
            r"C:\Windows\Fonts\times.ttf", r"C:\Windows\Fonts\timesbd.ttf",
            r"C:\Windows\Fonts\timesi.ttf", r"C:\Windows\Fonts\timesbi.ttf",
        ]),
        ("Tinos", [
            "/usr/share/fonts/truetype/tinos/Tinos-Regular.ttf",
            "/usr/share/fonts/truetype/tinos/Tinos-Bold.ttf",
            "/usr/share/fonts/truetype/tinos/Tinos-Italic.ttf",
            "/usr/share/fonts/truetype/tinos/Tinos-BoldItalic.ttf",
        ]),
        ("Liberation Serif", [
            "/usr/share/fonts/truetype/liberation2/LiberationSerif-Regular.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSerif-Bold.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSerif-Italic.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSerif-BoldItalic.ttf",
        ]),
        ("DejaVu Serif", [
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Italic.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSerif-BoldItalic.ttf",
        ]),
    ]


def _fc_match_file(family: str, style: str = "Regular", require_family: bool = True) -> Optional[str]:
    try:
        query = f"{family}:style={style}"
        result = subprocess.check_output(["fc-match", "-f", "%{family}|%{file}", query], text=True, stderr=subprocess.DEVNULL).strip()
        actual_family, _, path = result.partition("|")
        if require_family:
            wanted = family.casefold().replace(" ", "")
            actuals = [x.strip().casefold().replace(" ", "") for x in actual_family.split(",")]
            if wanted not in actuals:
                return None
        return path if path and os.path.isfile(path) else None
    except Exception:
        return None


def register_font_bundle() -> FontBundle:
    names = ["JournalSerif", "JournalSerif-Bold", "JournalSerif-Italic", "JournalSerif-BoldItalic"]

    # On Windows prefer the real installed Times New Roman files.
    win_times = _candidate_font_files()[0]
    if all(os.path.isfile(x) for x in win_times[1]):
        for name, path in zip(names, win_times[1]):
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, path, subfontIndex=0))
        return FontBundle(*names, win_times[0])

    # On Linux/macOS prefer a Times-compatible Unicode family resolved by fontconfig.
    # Tinos is metrically compatible with Times New Roman and keeps all four styles in one family.
    families = ["Times New Roman", "Tinos", "Liberation Serif", "DejaVu Serif"]
    for family in families:
        paths = [
            _fc_match_file(family, "Regular"),
            _fc_match_file(family, "Bold"),
            _fc_match_file(family, "Italic"),
            _fc_match_file(family, "Bold Italic"),
        ]
        if paths[0]:
            paths = [p or paths[0] for p in paths]
            for name, path in zip(names, paths):
                if name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(name, path, subfontIndex=0))
            return FontBundle(*names, family)

    # Last-resort known Linux paths.
    for label, files in _candidate_font_files()[1:]:
        if all(os.path.isfile(x) for x in files):
            for name, path in zip(names, files):
                if name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(name, path, subfontIndex=0))
            return FontBundle(*names, label)

    raise RuntimeError("Не знайдено Unicode-шрифт Times New Roman/Tinos/Liberation Serif/DejaVu Serif.")


def _nonempty(value: Any) -> bool:
    return value is not None and str(value) != ""


def _norm(value: Any) -> str:
    return safe_str(value).strip().casefold()


def _find_header_row(ws, keyword: str, start: int = 1, end: Optional[int] = None) -> List[Tuple[int, int]]:
    out = []
    end = min(ws.max_row, end or ws.max_row)
    for r in range(start, end + 1):
        for c in range(1, min(ws.max_column, 20) + 1):
            text = _norm(ws.cell(r, c).value)
            if keyword.casefold() in text:
                out.append((r, c))
    return out


def _parse_subject_teacher(ws, display_file: str) -> Tuple[str, str]:
    subject = ""
    teacher = ""
    for r in range(1, min(ws.max_row, 8) + 1):
        for c in range(1, min(ws.max_column, 12) + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            if "викладач" in v.casefold():
                parts = re.split(r"(?i)викладач\s*:?\s*", v, maxsplit=1)
                if len(parts) == 2:
                    before, after = parts
                    subject = before.strip(" \n\r\t:-")
                    teacher = after.strip()
                else:
                    teacher = v.strip()
                if not subject:
                    # nearest non-empty cell before/above
                    for rr in range(r, max(0, r - 2), -1):
                        for cc in range(c, 0, -1):
                            vv = ws.cell(rr, cc).value
                            if isinstance(vv, str) and vv.strip() and "викладач" not in vv.casefold():
                                subject = vv.strip()
                                break
                        if subject:
                            break
                break
        if teacher:
            break
    if not subject:
        subject = Path(display_file).stem
    if not teacher:
        teacher = "Викладач: не визначено"
    elif not teacher.casefold().startswith("викладач"):
        teacher = f"Викладач: {teacher}"
    return subject.strip(), teacher.strip()


def _detect_student_columns(ws, grade_row: int, theme_row: int) -> Tuple[Optional[int], Optional[int], List[int]]:
    end = max(grade_row + 1, theme_row - 1)
    best: Tuple[int, Optional[int], Optional[int]] = (0, None, None)
    for idx_col in range(1, min(6, ws.max_column)):
        name_col = idx_col + 1
        score = 0
        rows_with_idx = []
        for r in range(grade_row + 1, end + 1):
            idx = ws.cell(r, idx_col).value
            name = ws.cell(r, name_col).value
            numeric_idx = isinstance(idx, (int, float)) and not isinstance(idx, bool)
            if numeric_idx:
                rows_with_idx.append(r)
                if isinstance(name, str) and name.strip():
                    score += 2
                else:
                    score += 1
        if score > best[0]:
            best = (score, idx_col, name_col)
    if best[1] is None:
        return None, None, []
    rows = []
    for r in range(grade_row + 1, end + 1):
        idx = ws.cell(r, best[1]).value
        if isinstance(idx, (int, float)) and not isinstance(idx, bool):
            rows.append(r)
    return best[1], best[2], rows


def _find_theme_header(ws, grade_row: Optional[int]) -> Tuple[Optional[int], Optional[int], Optional[int], Optional[int]]:
    start = (grade_row or 1) + 1
    for r in range(start, min(ws.max_row, 300) + 1):
        date_col = hours_col = text_col = None
        for c in range(1, min(ws.max_column, 20) + 1):
            s = _norm(ws.cell(r, c).value)
            if s == "дата" or s.startswith("дата"):
                date_col = date_col or c
            if "кількість" in s and "год" in s:
                hours_col = hours_col or c
            if "теми занять" in s or ("теми" in s and "занять" in s):
                text_col = text_col or c
        if text_col is not None:
            return r, date_col, hours_col, text_col
    return None, None, None, None


def _meaningful_rows(ws, start: int, columns: Sequence[int]) -> List[int]:
    last = None
    for r in range(start, ws.max_row + 1):
        if any(_nonempty(ws.cell(r, c).value) for c in columns if c and c <= ws.max_column):
            last = r
    if last is None:
        return []
    return [r for r in range(start, last + 1) if any(_nonempty(ws.cell(r, c).value) for c in columns if c and c <= ws.max_column)]


def analyze_workbook(path: str, display_file: str, group: str) -> Tuple[List[JournalSource], List[WarningItem]]:
    warnings: List[WarningItem] = []
    journals: List[JournalSource] = []
    try:
        wb_formula = load_workbook(path, data_only=False, read_only=False)
        wb_value = load_workbook(path, data_only=True, read_only=False)
    except Exception as e:
        warnings.append(WarningItem(
            id=str(uuid.uuid4()), kind="workbook", message=f"Не вдалося прочитати файл: {e}",
            subject=Path(display_file).stem, correctable=False
        ))
        return journals, warnings

    for ws in wb_formula.worksheets:
        wsv = wb_value[ws.title]
        # Ignore empty/helper sheets.
        grade_candidates = _find_header_row(ws, "дата", 1, 20)
        grade_row = grade_candidates[0][0] if grade_candidates else None
        grade_date_col = grade_candidates[0][1] if grade_candidates else None
        theme_row, theme_date_col, theme_hours_col, theme_text_col = _find_theme_header(ws, grade_row)
        if grade_row is None and theme_row is None:
            continue

        subject, teacher = _parse_subject_teacher(ws, display_file)
        student_idx_col = student_name_col = None
        student_rows: List[int] = []
        grade_cols: List[int] = []
        note_rows: List[int] = []

        if grade_row is not None and theme_row is not None:
            student_idx_col, student_name_col, student_rows = _detect_student_columns(ws, grade_row, theme_row)
            if not student_rows:
                warnings.append(WarningItem(
                    id=str(uuid.uuid4()), kind="structure", message="Не вдалося надійно визначити список студентів.",
                    subject=subject, sheet=ws.title, correctable=False
                ))
            # Keep every non-empty marks column from after the 'Дата' stub through the last populated grade/notes cell.
            relevant_rows = [grade_row] + student_rows
            for r in range((max(student_rows) + 1) if student_rows else grade_row + 1, theme_row):
                if any(_nonempty(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 120) + 1)):
                    note_rows.append(r)
                    relevant_rows.append(r)
            if grade_date_col:
                last_col = grade_date_col
                for c in range(grade_date_col + 1, min(ws.max_column, 160) + 1):
                    if any(_nonempty(ws.cell(r, c).value) for r in relevant_rows):
                        last_col = c
                for c in range(grade_date_col + 1, last_col + 1):
                    # Drop only columns that contain absolutely no data in the relevant area.
                    if any(_nonempty(ws.cell(r, c).value) for r in relevant_rows):
                        grade_cols.append(c)

        theme_prefix_cols: List[int] = []
        theme_rows: List[int] = []
        if theme_row and theme_text_col:
            core_cols = [x for x in (theme_date_col, theme_hours_col, theme_text_col) if x]
            if theme_date_col:
                for c in range(1, theme_date_col):
                    if any(_nonempty(ws.cell(r, c).value) for r in range(theme_row + 1, ws.max_row + 1)):
                        theme_prefix_cols.append(c)
            theme_rows = _meaningful_rows(ws, theme_row + 1, theme_prefix_cols + core_cols)

        journal = JournalSource(
            file_path=path, display_file=display_file, sheet=ws.title,
            subject=subject, teacher=teacher, group=group,
            grade_header_row=grade_row, grade_date_col=grade_date_col,
            student_index_col=student_idx_col, student_name_col=student_name_col,
            student_rows=student_rows, grade_cols=grade_cols, note_rows=note_rows,
            theme_header_row=theme_row, theme_date_col=theme_date_col,
            theme_hours_col=theme_hours_col, theme_text_col=theme_text_col,
            theme_prefix_cols=theme_prefix_cols, theme_rows=theme_rows,
        )
        journals.append(journal)

        # Scan only cells that will be rendered for formula/error warnings.
        scan_cells: List[Cell] = []
        if grade_row and grade_date_col:
            for c in grade_cols:
                scan_cells.append(ws.cell(grade_row, c))
                for r in student_rows + note_rows:
                    scan_cells.append(ws.cell(r, c))
            if student_idx_col and student_name_col:
                for r in student_rows:
                    scan_cells.extend([ws.cell(r, student_idx_col), ws.cell(r, student_name_col)])
        if theme_row and theme_text_col:
            tcols = theme_prefix_cols + [x for x in (theme_date_col, theme_hours_col, theme_text_col) if x]
            for r in theme_rows:
                for c in tcols:
                    scan_cells.append(ws.cell(r, c))

        seen = set()
        for cell in scan_cells:
            if cell.coordinate in seen:
                continue
            seen.add(cell.coordinate)
            vcell = wsv[cell.coordinate]
            if cell.data_type == "f":
                cached = vcell.value
                if cached is None or is_excel_error(cached):
                    warnings.append(WarningItem(
                        id=f"{display_file}|{ws.title}|{cell.coordinate}",
                        kind="formula_error",
                        message="Формула не має коректного відображуваного результату. Можна вручну вказати значення лише для PDF.",
                        subject=subject, sheet=ws.title, cell=cell.coordinate,
                        original=safe_str(cached) if cached is not None else safe_str(cell.value),
                        formula=safe_str(cell.value), correctable=True,
                    ))
            elif cell.data_type == "e" or is_excel_error(cell.value):
                warnings.append(WarningItem(
                    id=f"{display_file}|{ws.title}|{cell.coordinate}",
                    kind="cell_error",
                    message="У клітинці Excel міститься помилка. Можна вручну вказати значення лише для PDF.",
                    subject=subject, sheet=ws.title, cell=cell.coordinate,
                    original=safe_str(cell.value), correctable=True,
                ))

    return journals, warnings


def safe_extract_xlsx(zip_path: str, out_dir: str) -> Tuple[List[Tuple[str, str]], str]:
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    extracted: List[Tuple[str, str]] = []
    group = ""
    with zipfile.ZipFile(zip_path) as zf:
        infos = [i for i in zf.infolist() if not i.is_dir()]
        if len(infos) > 500:
            raise ValueError("Архів містить понад 500 файлів.")
        total = sum(i.file_size for i in infos)
        if total > 500 * 1024 * 1024:
            raise ValueError("Розпакований архів перевищує 500 МБ.")
        for info in infos:
            raw_name = info.filename.replace("\\", "/")
            parts = [p for p in raw_name.split("/") if p not in ("", ".")]
            if any(p == ".." for p in parts):
                raise ValueError("Небезпечний шлях усередині ZIP.")
            display_parts = [decode_google_unicode_name(p) for p in parts]
            if not group and len(display_parts) > 1:
                group = display_parts[0]
            ext = Path(display_parts[-1]).suffix.lower() if display_parts else ""
            if ext not in {".xlsx", ".xlsm"}:
                continue
            if info.file_size > 50 * 1024 * 1024:
                raise ValueError(f"Файл {display_parts[-1]} перевищує 50 МБ.")
            dest = Path(out_dir) / f"{len(extracted)+1:03d}_{display_parts[-1]}"
            with zf.open(info) as src, open(dest, "wb") as dst:
                shutil.copyfileobj(src, dst)
            extracted.append((str(dest), "/".join(display_parts)))
    if not extracted:
        raise ValueError("В архіві не знайдено .xlsx/.xlsm файлів.")
    return extracted, group or "Група"


def analyze_archive(zip_path: str, work_dir: str, progress_cb=None) -> Dict[str, Any]:
    extract_dir = os.path.join(work_dir, "xlsx")
    files, group = safe_extract_xlsx(zip_path, extract_dir)
    all_journals: List[JournalSource] = []
    all_warnings: List[WarningItem] = []
    total_files = len(files)
    if progress_cb:
        progress_cb(0, total_files, "", "analysis")
    for file_index, (path, display) in enumerate(files, start=1):
        if progress_cb:
            progress_cb(file_index - 1, total_files, display, "analysis")
        journals, warnings = analyze_workbook(path, display, group)
        all_journals.extend(journals)
        all_warnings.extend(warnings)
        if progress_cb:
            progress_cb(file_index, total_files, display, "analysis")
    # Alphabetical by subject; stable tie-breakers make output deterministic.
    all_journals.sort(key=lambda j: (j.subject.casefold(), j.teacher.casefold(), j.display_file.casefold(), j.sheet.casefold()))
    payload = {
        "group": group,
        "files_count": len(files),
        "journals_count": len(all_journals),
        "subjects": [j.subject for j in all_journals],
        "warnings": [asdict(w) for w in all_warnings],
        "journals": [asdict(j) for j in all_journals],
    }
    with open(os.path.join(work_dir, "analysis.json"), "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    return payload


def load_analysis(work_dir: str) -> Dict[str, Any]:
    with open(os.path.join(work_dir, "analysis.json"), encoding="utf-8") as f:
        return json.load(f)


def _cell_text(ws_f, ws_v, r: int, c: int, display_file: str, sheet: str, overrides: Dict[str, Dict[str, Any]]) -> str:
    key = f"{display_file}|{sheet}|{ws_f.cell(r,c).coordinate}"
    override = overrides.get(key)
    if override:
        mode = override.get("mode", "keep")
        if mode == "replace":
            return str(override.get("value", ""))
        if mode == "blank":
            return ""
    return visible_cell_value(ws_f.cell(r, c), ws_v.cell(r, c))


def _draw_subject_header(c: canvas.Canvas, subject: str, teacher: str, fonts: FontBundle):
    c.setFillColor(colors.black)
    c.setStrokeColor(colors.black)
    c.setFont(fonts.bolditalic, 11)
    c.drawCentredString(PAGE_W / 2, PAGE_H - TOP - 2 * mm, subject)
    c.setFont(fonts.bolditalic, 10)
    c.drawCentredString(PAGE_W / 2, PAGE_H - TOP - 7 * mm, teacher)


def _measure_mark_col_width(texts: Iterable[str], fonts: FontBundle) -> float:
    width = 8.5 * mm
    for text in texts:
        s = str(text)
        if len(s) <= 3:
            continue
        width = max(width, min(14 * mm, pdfmetrics.stringWidth(s, fonts.regular, 5.8) + 2.2 * mm))
    return width


def _marks_chunks(j: JournalSource, ws_f, ws_v, fonts: FontBundle, overrides: Dict[str, Dict[str, Any]]) -> List[List[int]]:
    available = PAGE_W - LEFT - RIGHT
    idx_w = 9 * mm
    name_w = 70 * mm
    fixed = idx_w + name_w
    chunks: List[List[int]] = []
    current: List[int] = []
    used = fixed
    for c in j.grade_cols:
        texts = [_cell_text(ws_f, ws_v, j.grade_header_row, c, j.display_file, j.sheet, overrides)]
        texts += [_cell_text(ws_f, ws_v, r, c, j.display_file, j.sheet, overrides) for r in j.student_rows]
        col_w = _measure_mark_col_width(texts, fonts)
        # Most grade columns look best narrow; long values get a little more room.
        col_w = max(8.5 * mm, min(col_w, 14 * mm))
        if current and used + col_w > available:
            chunks.append(current)
            current = []
            used = fixed
        current.append(c)
        used += col_w
    if current:
        chunks.append(current)
    return chunks or [[]]


def _render_marks_pdf(j: JournalSource, out_path: str, fonts: FontBundle, overrides: Dict[str, Dict[str, Any]]):
    wb_f = load_workbook(j.file_path, data_only=False, read_only=False)
    wb_v = load_workbook(j.file_path, data_only=True, read_only=False)
    ws_f = wb_f[j.sheet]; ws_v = wb_v[j.sheet]
    c = canvas.Canvas(out_path, pagesize=PAGE_SIZE, pageCompression=1, initialFontName=fonts.regular, initialFontSize=8)
    chunks = _marks_chunks(j, ws_f, ws_v, fonts, overrides)

    for chunk in chunks:
        _draw_subject_header(c, j.subject, j.teacher, fonts)
        x0 = LEFT
        top_y = PAGE_H - TOP - HEADER_H
        idx_w = 9 * mm
        name_w = 70 * mm
        available = PAGE_W - LEFT - RIGHT
        remaining = max(1, len(chunk))
        grade_w = (available - idx_w - name_w) / remaining if remaining else 0
        # Dynamic header height for rotated date/assessment labels.
        header_texts = [_cell_text(ws_f, ws_v, j.grade_header_row, gc, j.display_file, j.sheet, overrides) for gc in chunk]
        max_header_text_w = max([pdfmetrics.stringWidth(t, fonts.regular, 7) for t in header_texts] + [0])
        header_h = max(30 * mm, min(46 * mm, max_header_text_w + 5 * mm))
        row_count = max(1, len(j.student_rows) + len(j.note_rows))
        usable_h = top_y - BOTTOM - header_h
        row_h = min(8 * mm, usable_h / row_count)
        font_size = min(8.2, max(5.4, row_h / mm * 1.35))

        # Table data built with Platypus, drawn onto the canvas for consistent wrapping.
        data: List[List[Any]] = []
        header = [PlainCellText("Дата", fonts.bold, 9, "center"), PlainCellText("", fonts.bold, 8, "center")]
        header.extend([VerticalText(t, fonts.italic, 7) for t in header_texts])
        data.append(header)

        for r in j.student_rows:
            idx = _cell_text(ws_f, ws_v, r, j.student_index_col, j.display_file, j.sheet, overrides) if j.student_index_col else ""
            name = _cell_text(ws_f, ws_v, r, j.student_name_col, j.display_file, j.sheet, overrides) if j.student_name_col else ""
            row: List[Any] = [PlainCellText(idx, fonts.regular, font_size, "center"), PlainCellText(name, fonts.italic, font_size, "left")]
            row += [PlainCellText(_cell_text(ws_f, ws_v, r, gc, j.display_file, j.sheet, overrides), fonts.regular, font_size, "center") for gc in chunk]
            data.append(row)

        for r in j.note_rows:
            prefix_vals = []
            for cc in range(1, (j.grade_date_col or 1) + 1):
                txt = _cell_text(ws_f, ws_v, r, cc, j.display_file, j.sheet, overrides)
                if txt != "": prefix_vals.append(txt)
            label = " ".join(prefix_vals)
            row = [PlainCellText("", fonts.regular, max(5.2, font_size - 0.8), "center"), PlainCellText(label, fonts.italic, max(5.2, font_size - 0.8), "left")]
            row += [PlainCellText(_cell_text(ws_f, ws_v, r, gc, j.display_file, j.sheet, overrides), fonts.regular, max(5.0, font_size - 1), "center") for gc in chunk]
            data.append(row)

        col_widths = [idx_w, name_w] + [grade_w] * len(chunk)
        row_heights = [header_h] + [row_h] * (len(data) - 1)
        table = Table(data, colWidths=col_widths, rowHeights=row_heights)
        table.setStyle(TableStyle([
            ("GRID", (0,0), (-1,-1), 0.65, colors.black),
            ("FONTNAME", (0,0), (-1,-1), fonts.regular),
            ("FONTSIZE", (0,0), (-1,-1), 8),
            ("SPAN", (0,0), (1,0)),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("LEFTPADDING", (0,0), (-1,-1), 1.5),
            ("RIGHTPADDING", (0,0), (-1,-1), 1.5),
            ("TOPPADDING", (0,0), (-1,-1), 1.0),
            ("BOTTOMPADDING", (0,0), (-1,-1), 1.0),
        ]))
        tw, th = table.wrapOn(c, available, top_y - BOTTOM)
        table.drawOn(c, LEFT, top_y - th)
        c.showPage()
    c.save()


def _topic_header_on_page(subject: str, teacher: str, fonts: FontBundle):
    def fn(c: canvas.Canvas, doc):
        _draw_subject_header(c, subject, teacher, fonts)
    return fn


def _render_topics_pdf(j: JournalSource, out_path: str, fonts: FontBundle, overrides: Dict[str, Dict[str, Any]]):
    wb_f = load_workbook(j.file_path, data_only=False, read_only=False)
    wb_v = load_workbook(j.file_path, data_only=True, read_only=False)
    ws_f = wb_f[j.sheet]; ws_v = wb_v[j.sheet]

    doc = SimpleDocTemplate(
        out_path, pagesize=PAGE_SIZE,
        leftMargin=LEFT, rightMargin=RIGHT,
        topMargin=TOP + HEADER_H, bottomMargin=BOTTOM,
        title=j.subject, author="Journal PDF Web"
    )
    story: List[Any] = []
    columns: List[Tuple[str, int]] = []
    for pc in j.theme_prefix_cols:
        hdr = _cell_text(ws_f, ws_v, j.theme_header_row, pc, j.display_file, j.sheet, overrides)
        columns.append((hdr or "№", pc))
    if j.theme_date_col:
        columns.append((_cell_text(ws_f, ws_v, j.theme_header_row, j.theme_date_col, j.display_file, j.sheet, overrides) or "Дата", j.theme_date_col))
    if j.theme_hours_col:
        columns.append((_cell_text(ws_f, ws_v, j.theme_header_row, j.theme_hours_col, j.display_file, j.sheet, overrides) or "Кількість годин", j.theme_hours_col))
    if j.theme_text_col:
        columns.append((_cell_text(ws_f, ws_v, j.theme_header_row, j.theme_text_col, j.display_file, j.sheet, overrides) or "Теми занять", j.theme_text_col))

    data: List[List[Any]] = []
    data.append([PlainCellText(h, fonts.bold, 8, "center") for h, _ in columns])
    for r in j.theme_rows:
        row: List[Any] = []
        for idx, (_, col) in enumerate(columns):
            txt = _cell_text(ws_f, ws_v, r, col, j.display_file, j.sheet, overrides)
            align = "left" if col == j.theme_text_col else "center"
            fs = 7.8 if col == j.theme_text_col else 7.5
            row.append(PlainCellText(txt, fonts.regular if col == j.theme_text_col else fonts.italic, fs, align))
        data.append(row)

    available = PAGE_W - LEFT - RIGHT
    prefix_n = len(j.theme_prefix_cols)
    widths: List[float] = []
    if prefix_n:
        widths += [10 * mm] * prefix_n
    if j.theme_date_col:
        widths.append(20 * mm)
    if j.theme_hours_col:
        widths.append(25 * mm)
    used = sum(widths)
    if j.theme_text_col:
        widths.append(max(60 * mm, available - used))

    table = Table(data, colWidths=widths, repeatRows=1, splitByRow=1)
    table.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.65, colors.black),
        ("FONTNAME", (0,0), (-1,-1), fonts.regular),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 2.0),
        ("RIGHTPADDING", (0,0), (-1,-1), 2.0),
        ("TOPPADDING", (0,0), (-1,-1), 2.0),
        ("BOTTOMPADDING", (0,0), (-1,-1), 2.0),
    ]))
    story.append(table)
    def _canvasmaker(filename, **kwargs):
        if not kwargs.get("initialFontName"):
            kwargs["initialFontName"] = fonts.regular
        if not kwargs.get("initialFontSize"):
            kwargs["initialFontSize"] = 8
        return canvas.Canvas(filename, **kwargs)
    doc.build(
        story,
        onFirstPage=_topic_header_on_page(j.subject, j.teacher, fonts),
        onLaterPages=_topic_header_on_page(j.subject, j.teacher, fonts),
        canvasmaker=_canvasmaker,
    )


def generate_pdf(work_dir: str, overrides: Dict[str, Dict[str, Any]], output_path: str, progress_cb=None) -> Dict[str, Any]:
    analysis = load_analysis(work_dir)
    fonts = register_font_bundle()
    temp_pdfs: List[str] = []
    temp_dir = os.path.join(work_dir, "pdf_parts")
    os.makedirs(temp_dir, exist_ok=True)

    total_journals = len(analysis["journals"])
    if progress_cb:
        progress_cb(0, total_journals, "", "pdf")
    for idx, jdata in enumerate(analysis["journals"], start=1):
        j = JournalSource(**jdata)
        if progress_cb:
            progress_cb(idx - 1, total_journals, j.subject, "pdf")
        if j.grade_header_row and j.student_rows and j.grade_cols:
            p = os.path.join(temp_dir, f"{idx:03d}_marks.pdf")
            _render_marks_pdf(j, p, fonts, overrides)
            temp_pdfs.append(p)
        if j.theme_header_row and j.theme_text_col and j.theme_rows:
            p = os.path.join(temp_dir, f"{idx:03d}_topics.pdf")
            _render_topics_pdf(j, p, fonts, overrides)
            temp_pdfs.append(p)
        if progress_cb:
            progress_cb(idx, total_journals, j.subject, "pdf")

    if not temp_pdfs:
        raise RuntimeError("Не вдалося сформувати жодної сторінки PDF.")

    writer = PdfWriter()
    writer.add_metadata({
        "/Title": f"Журнал групи {analysis.get('group','')}",
        "/Author": "Journal PDF Web",
        "/Subject": "Архівний журнал групи",
        "/Keywords": "журнал, група, електронний журнал, архів",
    })
    for p in temp_pdfs:
        reader = PdfReader(p)
        for page in reader.pages:
            writer.add_page(page)
    with open(output_path, "wb") as f:
        writer.write(f)

    return {"font": fonts.label, "pages": len(PdfReader(output_path).pages), "group": analysis.get("group", "")}
